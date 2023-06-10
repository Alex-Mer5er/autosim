from docx import Document
from app.browser import PersonaForm

class DocTable():
    def __init__(self, path_to_docx):
        self.doc = Document(path_to_docx)
        self.table = self.doc.tables[0]
        self.columns = {}
        self.search_columns()

    def search_columns(self):
        #columns = ['Фамилия', 'Имя', 'Отчество', 'Дата рождения', 'Серия паспорта', 'Номер паспорта', 'Дата выдачи паспорта', 'Кем выдан паспорт']
        key_columns = {'Фамилия': 'familiya', 'Имя': 'imya', 'Отчество': 'otchestvo', 'Дата рождения': 'dr', 'Серия паспорта': 'seriya', 'Номер паспорта': 'nomer', 'Дата выдачи паспорта': 'dv', 'Кем выдан паспорт': 'vydan'}
        for i, cell in enumerate(self.table.rows[0].cells):
            if cell.text in key_columns:
                self.columns[key_columns[cell.text]] = i
        assert len(self.columns.keys()) == 8

    def search_person(self, s, n):
        for i, cs in enumerate(self.table.columns[self.columns['seriya']].cells):
            if cs.text == s:
                if self.table.cell(i, self.columns['nomer']).text == n:
                    return PersonaForm(familiya=self.table.cell(i, self.columns['familiya']).text, imya=self.table.cell(i, self.columns['imya']).text, otchestvo=self.table.cell(i, self.columns['otchestvo']).text, dr=self.table.cell(i, self.columns['dr']).text, seriya=self.table.cell(i, self.columns['seriya']).text, nomer=self.table.cell(i, self.columns['nomer']).text, vydan=self.table.cell(i, self.columns['vydan']).text, dv=self.table.cell(i, self.columns['dv']).text)
        return None
